from openpyxl import load_workbook
import os

class ExcelUtil:

    @staticmethod
    def get_test_data(sheet):
        project_root = os.path.dirname(os.path.dirname(__file__))
        excel_path = os.path.join(project_root, "data", "data.xlsx")

        print("Using Excel file:", excel_path)

        if not os.path.exists(excel_path):
            raise FileNotFoundError(f"Excel file not found at {excel_path}")

        wb = load_workbook(excel_path)
        ws = wb[sheet]

        data = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            data.append(row)
        return data



# from openpyxl import load_workbook
#
# class ExcelUtil:
#
#     @staticmethod
#     def get_data(sheet, row, col, path="data.xlsx"):
#         wb = load_workbook(path)
#         return wb[sheet].cell(row=row, column=col).value
